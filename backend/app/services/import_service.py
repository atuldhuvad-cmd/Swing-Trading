import csv
import io
import hashlib
from datetime import datetime
from typing import List, Dict, Any, Tuple
from sqlalchemy.orm import Session
import json

from app.models import (
    ImportBatch, ImportBatchDetail, ReviewQueue, 
    StockMaster, BrokerMaster, BrokerAlias, RatingNormalization,
    BrokerRecommendation, SourceReference, RecommendationSource,
    SourceTypeMaster
)
from app.schemas.import_batch import ColumnMapping, ParsedImportRow

# Optional openpyxl for xlsx
try:
    import openpyxl
except ImportError:
    openpyxl = None

class ImportService:
    @staticmethod
    def parse_file(file_content: bytes, filename: str) -> Tuple[List[str], List[Dict[str, Any]]]:
        if filename.endswith('.csv'):
            text_content = file_content.decode('utf-8-sig') # Handle BOM
            reader = csv.reader(io.StringIO(text_content))
            headers = next(reader, [])
            headers = [h.strip() for h in headers]
            rows = []
            for row in reader:
                if any(row): # Skip empty rows
                    rows.append(dict(zip(headers, row)))
            return headers, rows
        elif filename.endswith('.xlsx'):
            if not openpyxl:
                raise ValueError("openpyxl is not installed")
            wb = openpyxl.load_workbook(filename=io.BytesIO(file_content), read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            headers = next(rows_iter, [])
            headers = [str(h).strip() if h else "" for h in headers]
            rows = []
            for row in rows_iter:
                if any(row):
                    rows.append(dict(zip(headers, row)))
            return headers, rows
        else:
            raise ValueError("Unsupported file format. Only .csv and .xlsx are supported.")

    @staticmethod
    def _normalize_string(val: Any) -> str | None:
        if val is None:
            return None
        val_str = str(val).strip()
        if val_str.lower() in ('', 'na', 'n/a', 'not available', 'null', 'none'):
            return None
        return val_str

    @staticmethod
    def _parse_price(val: Any) -> float | None:
        val_str = ImportService._normalize_string(val)
        if not val_str:
            return None
        # Remove currency symbols and commas
        val_str = val_str.replace('₹', '').replace(',', '').replace('Rs.', '').replace('Rs', '').strip()
        try:
            return float(val_str)
        except ValueError:
            return None

    @staticmethod
    def _parse_date(val: Any) -> datetime | None:
        if isinstance(val, datetime):
            return val
        val_str = ImportService._normalize_string(val)
        if not val_str:
            return None
        
        # Some sources (e.g. a compiled "current calls" page) report a row
        # as being current as of today rather than giving a specific date.
        # Treat these known placeholders as today's date instead of
        # rejecting the row outright.
        if val_str.lower() in ('current note', 'current', 'today', 'as of today'):
            return datetime.utcnow()

        # Try common formats
        formats = [
            '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%m/%d/%Y',
            '%Y-%m-%dT%H:%M:%S', '%Y-%m-%d %H:%M:%S',
            '%d %b %Y', '%d %B %Y',
        ]
        for fmt in formats:
            try:
                dt = datetime.strptime(val_str, fmt)
                return dt
            except ValueError:
                pass
        return None

    @staticmethod
    def normalize_row(db: Session, raw_row: Dict[str, Any], mapping: ColumnMapping) -> Tuple[Dict[str, Any], str, str]:
        """Returns (mapped_data, action, error_message)"""
        mapped = {}
        error_msg = None
        action = "UNIQUE"
        
        # Map fields
        def get_val(field: str) -> Any:
            mapped_header = getattr(mapping, field)
            if mapped_header and mapped_header in raw_row:
                return raw_row[mapped_header]
            return None

        # 1. Stock Resolution
        nse_symbol_raw = ImportService._normalize_string(get_val('nse_symbol'))
        if not nse_symbol_raw:
            return mapped, "INVALID", "Missing NSE Symbol"
            
        nse_symbol = nse_symbol_raw.upper().strip()
        stock = db.query(StockMaster).filter(StockMaster.nse_symbol == nse_symbol).first()
        if not stock:
            action = "UNKNOWN_STOCK"
            error_msg = f"Stock {nse_symbol} not found in master"
        
        mapped['nse_symbol'] = nse_symbol
        mapped['stock_id'] = stock.stock_id if stock else None

        # 2. Broker Resolution
        broker_name_raw = ImportService._normalize_string(get_val('broker_name'))
        if not broker_name_raw:
            return mapped, "INVALID", "Missing Broker Name"
            
        norm_broker_name = broker_name_raw.lower().replace(' ', '')
        # Try canonical/normalized
        broker = db.query(BrokerMaster).filter(BrokerMaster.normalized_name == norm_broker_name).first()
        if not broker:
            # Try display name
            broker = db.query(BrokerMaster).filter(BrokerMaster.display_name.ilike(f"%{broker_name_raw}%")).first()
            if not broker:
                # Try alias
                alias = db.query(BrokerAlias).filter(BrokerAlias.alias_name.ilike(f"%{broker_name_raw}%")).first()
                if alias:
                    broker = alias.broker
        
        if not broker:
            if action == "UNIQUE":
                action = "UNKNOWN_BROKER"
                error_msg = f"Broker {broker_name_raw} not found"
            else:
                action = "INVALID"
                error_msg = f"Unknown Stock and Broker"
        
        mapped['broker_name'] = broker.display_name if broker else broker_name_raw
        mapped['broker_id'] = broker.broker_id if broker else None

        # 3. Date
        rec_date = ImportService._parse_date(get_val('recommendation_date'))
        if not rec_date:
            return mapped, "INVALID", "Missing or invalid recommendation date"
        if rec_date > datetime.utcnow():
            return mapped, "INVALID", "Future recommendation date not allowed"
        mapped['recommendation_date'] = rec_date.isoformat()

        # 4. Rating
        orig_rating = ImportService._normalize_string(get_val('original_rating'))
        if not orig_rating:
            return mapped, "INVALID", "Missing original rating"
        mapped['original_rating'] = orig_rating
        
        # Try mapping rating
        norm_rating_db = db.query(RatingNormalization).filter(RatingNormalization.original_rating.ilike(orig_rating)).first()
        mapped['normalized_rating'] = norm_rating_db.normalized_rating if norm_rating_db else None
        
        if not norm_rating_db and action in ("UNIQUE", "UNKNOWN_STOCK", "UNKNOWN_BROKER"):
            action = "REVIEW_REQUIRED"
            error_msg = f"Unknown rating {orig_rating}"

        # 5. Prices
        target = ImportService._parse_price(get_val('target_price'))
        if target is not None and target <= 0:
            return mapped, "INVALID", "Target price must be > 0"
        mapped['target_price'] = target

        mapped['recommended_price'] = ImportService._parse_price(get_val('recommended_price'))
        
        low = ImportService._parse_price(get_val('entry_price_low'))
        high = ImportService._parse_price(get_val('entry_price_high'))
        if low is not None and high is not None and low > high:
            return mapped, "INVALID", "Entry low cannot be greater than entry high"
        mapped['entry_price_low'] = low
        mapped['entry_price_high'] = high
        mapped['stop_loss'] = ImportService._parse_price(get_val('stop_loss'))
        
        # 6. Source
        mapped['source_name'] = ImportService._normalize_string(get_val('source_name'))
        mapped['source_url'] = ImportService._normalize_string(get_val('source_url'))

        source_date_raw = ImportService._parse_date(get_val('source_date'))
        mapped['source_date'] = (
            source_date_raw.isoformat() if source_date_raw else mapped.get('recommendation_date')
        )

        original_text = ImportService._normalize_string(get_val('original_text'))
        if not original_text:
            # No free-text column in the file -- synthesize a short, factual
            # summary from the row itself so the evidence trail isn't empty.
            summary_bits = [mapped.get('broker_name') or broker_name_raw, orig_rating]
            if mapped.get('target_price'):
                summary_bits.append(f"target Rs {mapped['target_price']}")
            if mapped.get('recommended_price'):
                summary_bits.append(f"CMP Rs {mapped['recommended_price']}")
            if mapped.get('source_name'):
                summary_bits.append(f"via {mapped['source_name']}")
            original_text = f"{nse_symbol}: " + ", ".join(str(b) for b in summary_bits if b)
        mapped['original_text'] = original_text

        # Source Type resolution -- a per-row column takes priority; if the
        # file doesn't have one (e.g. a compiled file that is entirely one
        # kind of source), fall back to the batch-level default supplied in
        # the mapping.
        source_type_raw = ImportService._normalize_string(get_val('source_type'))
        if not source_type_raw:
            source_type_raw = ImportService._normalize_string(getattr(mapping, 'default_source_type', None))
        if source_type_raw:
            st_master = db.query(SourceTypeMaster).filter(SourceTypeMaster.type_name.ilike(source_type_raw)).first()
            if st_master:
                mapped['source_type_id'] = st_master.source_type_id
            else:
                if action in ("UNIQUE", "UNKNOWN_STOCK", "UNKNOWN_BROKER"):
                    action = "REVIEW_REQUIRED"
                    error_msg = f"Unknown source type: {source_type_raw}"
        else:
            if action in ("UNIQUE", "UNKNOWN_STOCK", "UNKNOWN_BROKER"):
                action = "REVIEW_REQUIRED"
                error_msg = "Missing source type"
                
        # Verification Status
        verif = ImportService._normalize_string(get_val('verification_status'))
        if verif not in ['VERIFIED_PRIMARY', 'VERIFIED_SECONDARY', 'REJECTED']:
            verif = 'PROVISIONAL'
        mapped['verification_status'] = verif

        return mapped, action, error_msg

    @staticmethod
    def fingerprint(mapped_data: Dict[str, Any]) -> str | None:
        stock_id = mapped_data.get('stock_id')
        broker_id = mapped_data.get('broker_id')
        date_str = mapped_data.get('recommendation_date')
        rating = mapped_data.get('normalized_rating')
        if not all([stock_id, broker_id, date_str, rating]):
            return None
            
        base_string = f"{stock_id}-{broker_id}-{date_str}-{rating}"
        return hashlib.sha256(base_string.encode()).hexdigest()

    @staticmethod
    def deduplicate(db: Session, mapped_data: Dict[str, Any], current_action: str, batch_fingerprints: set) -> Tuple[str, str, int | None]:
        if current_action not in ("UNIQUE", "REVIEW_REQUIRED"):
            return current_action, None, None
            
        fp = ImportService.fingerprint(mapped_data)
        if not fp:
            return current_action, None, None
            
        # 1. Check intra-batch
        if fp in batch_fingerprints:
            return "DUPLICATE_IN_BATCH", "Duplicate row in uploaded file", None
            
        batch_fingerprints.add(fp)
        
        # 2. Check Database Exact
        existing = db.query(BrokerRecommendation).filter(BrokerRecommendation.fingerprint == fp).first()
        if existing:
            # Exact Match. Check if URL is new.
            url = mapped_data.get('source_url')
            if url:
                has_source = db.query(SourceReference).join(RecommendationSource).filter(
                    RecommendationSource.recommendation_id == existing.recommendation_id,
                    SourceReference.url == url
                ).first()
                if not has_source:
                    return "ATTACH_SOURCE", "Exact duplicate, new source to attach", existing.recommendation_id
            return "EXACT_DUPLICATE", "Already exists in database", existing.recommendation_id

        # 3. Check Probable Duplicate / Possible Update
        # Same stock and broker, close date?
        rec_date = datetime.fromisoformat(mapped_data['recommendation_date'])
        stock_id = mapped_data['stock_id']
        broker_id = mapped_data['broker_id']
        
        # Check existing recs for same stock/broker
        recent_recs = db.query(BrokerRecommendation).filter(
            BrokerRecommendation.stock_id == stock_id,
            BrokerRecommendation.broker_id == broker_id
        ).all()
        
        for r in recent_recs:
            if abs((r.recommendation_date - rec_date).days) <= 7:
                if r.normalized_rating == mapped_data.get('normalized_rating'):
                    return "PROBABLE_DUPLICATE", f"Very similar to rec {r.recommendation_id}", r.recommendation_id
                else:
                    return "POSSIBLE_UPDATE", f"Possible update to rec {r.recommendation_id}", r.recommendation_id

        return current_action, None, None
    @staticmethod
    def confirm_batch(db: Session, batch_id: int):
        batch = db.query(ImportBatch).filter(ImportBatch.batch_id == batch_id).first()
        if not batch:
            raise ValueError(f'Invalid batch: None for id {batch_id}')
        if batch.status not in ('PREVIEW', 'COMPLETED'):
            raise ValueError(f'Invalid status: {batch.status} for id {batch_id}')
            
        details = db.query(ImportBatchDetail).filter(ImportBatchDetail.batch_id == batch_id).all()
        for d in details:
            if d.status != 'PREVIEW':
                continue
                
            if d.action == 'UNIQUE':
                mapped = json.loads(d.mapped_data)

                if not mapped.get('source_type_id'):
                    # Defensive: normalize_row should already have routed
                    # this to REVIEW_REQUIRED, but never insert a
                    # SourceReference with no source type -- the column is
                    # NOT NULL and would crash the whole confirm.
                    d.status = 'REJECTED'
                    d.error_message = 'Missing source type -- cannot create evidence record'
                    batch.rejected_rows += 1
                    continue

                # Create source
                src = SourceReference(
                    source_type_id=mapped.get('source_type_id'),
                    publication_name=mapped.get('source_name'),
                    url=mapped.get('source_url'),
                    source_date=datetime.fromisoformat(mapped['source_date']) if mapped.get('source_date') else None,
                    original_text=mapped.get('original_text'),
                    verification_status=mapped.get('verification_status', 'PROVISIONAL'),
                    import_batch_id=batch_id
                )
                db.add(src)
                db.flush()
                
                # Create rec
                rec = BrokerRecommendation(
                    stock_id=mapped['stock_id'],
                    broker_id=mapped['broker_id'],
                    recommendation_date=datetime.fromisoformat(mapped['recommendation_date']),
                    original_rating=mapped['original_rating'],
                    normalized_rating=mapped['normalized_rating'],
                    target_price=mapped.get('target_price'),
                    recommended_price=mapped.get('recommended_price'),
                    entry_price_low=mapped.get('entry_price_low'),
                    entry_price_high=mapped.get('entry_price_high'),
                    stop_loss=mapped.get('stop_loss'),
                    lifecycle_status='CURRENT',
                    import_batch_id=batch_id
                )
                rec.fingerprint = ImportService.fingerprint(mapped)
                db.add(rec)
                db.flush()
                
                # Link
                db.add(RecommendationSource(recommendation_id=rec.recommendation_id, source_reference_id=src.source_reference_id))
                
                d.status = 'COMPLETED'
                d.recommendation_id = rec.recommendation_id
                d.source_reference_id = src.source_reference_id
                batch.accepted_rows += 1
                
            elif d.action == 'ATTACH_SOURCE':
                mapped = json.loads(d.mapped_data)

                if not mapped.get('source_type_id'):
                    d.status = 'REJECTED'
                    d.error_message = 'Missing source type -- cannot create evidence record'
                    batch.rejected_rows += 1
                    continue

                src = SourceReference(
                    source_type_id=mapped.get('source_type_id'),
                    publication_name=mapped.get('source_name'),
                    url=mapped.get('source_url'),
                    source_date=datetime.fromisoformat(mapped['source_date']) if mapped.get('source_date') else None,
                    original_text=mapped.get('original_text'),
                    verification_status=mapped.get('verification_status', 'PROVISIONAL'),
                    import_batch_id=batch_id
                )
                db.add(src)
                db.flush()
                db.add(RecommendationSource(recommendation_id=d.recommendation_id, source_reference_id=src.source_reference_id))
                
                d.status = 'COMPLETED'
                d.source_reference_id = src.source_reference_id
                batch.accepted_rows += 1
                
            elif d.action in ('PROBABLE_DUPLICATE', 'POSSIBLE_UPDATE', 'REVIEW_REQUIRED', 'UNKNOWN_STOCK', 'UNKNOWN_BROKER'):
                rq = ReviewQueue(
                    item_type='IMPORT_ROW',
                    item_reference_id=d.detail_id,
                    reason=d.error_message or d.action,
                    import_batch_id=batch_id,
                    status='PENDING'
                )
                db.add(rq)
                d.status = 'REVIEW'
                batch.review_rows += 1
                
            elif d.action in ('EXACT_DUPLICATE', 'DUPLICATE_IN_BATCH'):
                d.status = 'IGNORED'
                batch.duplicate_rows += 1
                
            elif d.action == 'INVALID':
                d.status = 'REJECTED'
                batch.rejected_rows += 1
                
        batch.status = 'COMPLETED'
        db.commit()

    @staticmethod
    def rollback_batch(db: Session, batch_id: int):
        batch = db.query(ImportBatch).filter(ImportBatch.batch_id == batch_id).first()
        if not batch:
            raise ValueError('Batch not found')
            
        # Recommendations created by this batch
        batch_recs = db.query(BrokerRecommendation.recommendation_id).filter(BrokerRecommendation.import_batch_id == batch_id).all()
        batch_rec_ids = [r[0] for r in batch_recs]
        
        # Sources created by this batch
        batch_sources = db.query(SourceReference.source_reference_id).filter(SourceReference.import_batch_id == batch_id).all()
        batch_source_ids = [s[0] for s in batch_sources]
        
        # 1. Delete links to recommendations that are about to be deleted
        if batch_rec_ids:
            db.execute(RecommendationSource.__table__.delete().where(
                RecommendationSource.recommendation_id.in_(batch_rec_ids)
            ))
            
        # 2. Delete links created by ATTACH_SOURCE (new source -> existing rec)
        details = db.query(ImportBatchDetail).filter(ImportBatchDetail.batch_id == batch_id, ImportBatchDetail.status == 'COMPLETED').all()
        for d in details:
            if d.recommendation_id and d.source_reference_id and d.action == 'ATTACH_SOURCE':
                # Delete this specific link created by the batch
                db.execute(RecommendationSource.__table__.delete().where(
                    (RecommendationSource.recommendation_id == d.recommendation_id) &
                    (RecommendationSource.source_reference_id == d.source_reference_id)
                ))
                
        # 3. Now check which batch_source_ids are STILL linked to ANY recommendation.
        sources_to_delete = []
        for sid in batch_source_ids:
            links = db.query(RecommendationSource).filter(RecommendationSource.source_reference_id == sid).count()
            if links == 0:
                sources_to_delete.append(sid)
                
        # 4. Nullify foreign keys in ImportBatchDetail for this batch to prevent FK constraint failures on deletion
        if sources_to_delete or batch_rec_ids:
            db.execute(ImportBatchDetail.__table__.update().where(
                ImportBatchDetail.batch_id == batch_id
            ).values(recommendation_id=None, source_reference_id=None))
            
        # 5. Delete orphaned sources
        if sources_to_delete:
            db.query(SourceReference).filter(SourceReference.source_reference_id.in_(sources_to_delete)).delete(synchronize_session=False)
            
        # 6. Delete batch-created recommendations
        if batch_rec_ids:
            # First delete dependent status history
            db.execute(__import__('app.models').models.RecommendationStatusHistory.__table__.delete().where(
                __import__('app.models').models.RecommendationStatusHistory.recommendation_id.in_(batch_rec_ids)
            ))
            db.query(BrokerRecommendation).filter(BrokerRecommendation.recommendation_id.in_(batch_rec_ids)).delete(synchronize_session=False)
            
        # 6. Delete review queue items
        db.query(ReviewQueue).filter(ReviewQueue.import_batch_id == batch_id).delete(synchronize_session=False)
        
        # 7. Mark batch rolled back
        batch.status = 'ROLLED_BACK'
        batch.accepted_rows = 0
        batch.review_rows = 0
        
        db.commit()
